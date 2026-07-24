from __future__ import annotations

from copy import copy
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .enrichment import enrich_probe
from .models import ProbeDefinition, TriggerDefinition

HEADERS = {
    "type_policy": "Type de politique",
    "supervisor": "Type de superviseur",
    "policy_name": "Nom de la politique",
    "collection_method": "Méthode de Collecte",
    "environment": "Environnement",
    "target_property": "Propriété de la Cible",
    "resource_type": "Type de ressource",
    "category": "Catégorie",
    "lld": "LLD",
    "resource": "Ressource",
    "key": "Clef",
    "description": "Description",
    "frequency": "Fréquence",
    "preprocessing": "Prétraitement",
    "filters": "Filtres",
    "macros": "Macros",
    "trigger_name": "Triger Name",
    "condition": "Condition de l'alerte",
    "severity": "Sévérité",
    "alert_message": "Message de l'alerte",
    "retention": "Métrologie / rétention",
}

SEVERITY_LABELS = {
    "NOT_CLASSIFIED": "Non classé",
    "INFORMATION": "Information",
    "WARNING": "Avertissement",
    "AVERAGE": "Moyen",
    "HIGH": "Haut",
    "DISASTER": "Désastre",
}

SEVERITY_STYLES = {
    "NOT_CLASSIFIED": ("97AAB3", "FFFFFF"),
    "INFORMATION": ("7499FF", "FFFFFF"),
    "WARNING": ("FFC859", "000000"),
    "AVERAGE": ("FFA059", "000000"),
    "HIGH": ("E97659", "FFFFFF"),
    "DISASTER": ("E45959", "FFFFFF"),
}

DEPENDENCY_FONT = InlineFont(color="008000", b=True)
MACRO_PATTERN = re.compile(r"\{\$[^}]+\}")


def _normalise(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _find_header_row(sheet: Worksheet) -> tuple[int, dict[str, int]]:
    expected = {_normalise(value): key for key, value in HEADERS.items()}
    for row in range(1, min(sheet.max_row, 50) + 1):
        mapping: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            normalised = _normalise(sheet.cell(row, column).value)
            if normalised in expected:
                mapping[expected[normalised]] = column
        if "resource" in mapping and "key" in mapping and "description" in mapping:
            return row, mapping
    raise ValueError("The catalogue model does not contain the expected header row")


def _capture_row_style(sheet: Worksheet, source_row: int) -> tuple[float | None, list[dict[str, object]]]:
    styles: list[dict[str, object]] = []
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(source_row, column)
        styles.append({
            "style": copy(source._style),
            "number_format": source.number_format,
            "alignment": copy(source.alignment),
            "protection": copy(source.protection),
        })
    return sheet.row_dimensions[source_row].height, styles


def _apply_row_style(
    sheet: Worksheet,
    target_row: int,
    row_style: tuple[float | None, list[dict[str, object]]],
) -> None:
    height, styles = row_style
    sheet.row_dimensions[target_row].height = height
    for column, style in enumerate(styles, start=1):
        target = sheet.cell(target_row, column)
        target._style = copy(style["style"])
        target.number_format = str(style["number_format"])
        target.alignment = copy(style["alignment"])
        target.protection = copy(style["protection"])


def _set(sheet: Worksheet, row: int, columns: dict[str, int], key: str, value: object) -> None:
    column = columns.get(key)
    if column:
        sheet.cell(row, column).value = value


def _apply_trigger_style(
    sheet: Worksheet,
    row: int,
    columns: dict[str, int],
    trigger: TriggerDefinition,
) -> None:
    colors = SEVERITY_STYLES.get(trigger.severity.upper())
    if not colors:
        return

    fill_color, font_color = colors
    for key in ("trigger_name", "severity"):
        column = columns.get(key)
        if not column:
            continue
        cell = sheet.cell(row, column)
        cell.fill = PatternFill(fill_type="solid", fgColor=fill_color)
        cell.font = Font(
            name=cell.font.name,
            size=cell.font.size,
            bold=cell.font.bold,
            italic=cell.font.italic,
            vertAlign=cell.font.vertAlign,
            underline=cell.font.underline,
            strike=cell.font.strike,
            color=font_color,
        )


def _trigger_condition(trigger: TriggerDefinition) -> str:
    if trigger.expression and trigger.recovery_expression:
        return f"Déclenchement : {trigger.expression}\nRétablissement : {trigger.recovery_expression}"
    return trigger.expression


def _retention(probe: ProbeDefinition) -> str:
    values = []
    if probe.history:
        values.append(f"Historique : {probe.history}")
    if probe.trends:
        values.append(f"Tendances : {probe.trends}")
    return " / ".join(values)


def _resource_value(probe: ProbeDefinition) -> str | CellRichText:
    if not probe.dependency_names:
        return probe.name

    value = CellRichText()
    for dependency_name in probe.dependency_names:
        value.append(TextBlock(DEPENDENCY_FONT, f"Dépend de : {dependency_name}\n"))
    value.append(probe.name)
    return value


def _used_macros(probe: ProbeDefinition, trigger: TriggerDefinition | None) -> str:
    texts = [
        probe.name,
        probe.key,
        probe.description,
        probe.calculation_formula,
        *probe.preprocessing,
        *probe.filters,
    ]
    if trigger:
        texts.extend(
            [
                trigger.name,
                trigger.expression,
                trigger.recovery_expression,
                trigger.description,
            ]
        )

    macro_index = dict(probe.template_macros)
    found: list[str] = []
    for text in texts:
        for macro_name in MACRO_PATTERN.findall(text or ""):
            if macro_name in macro_index and macro_name not in found:
                found.append(macro_name)
    return "\n".join(f"{name} = {macro_index[name]}" for name in found)


def _row_values(probe: ProbeDefinition, trigger: TriggerDefinition | None) -> dict[str, object]:
    enriched = enrich_probe(probe)
    return {
        "type_policy": "STD",
        "supervisor": "Zabbix",
        "policy_name": probe.template_name,
        "collection_method": enriched.collection_method,
        "environment": "PROD/QUAL",
        "target_property": enriched.target_property,
        "resource_type": enriched.resource_type,
        "category": enriched.category,
        "lld": "Yes" if probe.lld else "No",
        "resource": _resource_value(probe),
        "key": probe.key,
        "description": probe.description,
        "frequency": probe.delay,
        "preprocessing": "\n".join(probe.preprocessing),
        "filters": "\n".join(probe.filters),
        "macros": _used_macros(probe, trigger),
        "trigger_name": trigger.name if trigger else "",
        "condition": _trigger_condition(trigger) if trigger else "",
        "severity": SEVERITY_LABELS.get(trigger.severity, trigger.severity) if trigger else "",
        "alert_message": trigger.description if trigger else "",
        "retention": _retention(probe),
    }


def _write_tags_sheet(workbook, probes: list[ProbeDefinition]) -> None:
    title = "Tags"
    sheet = workbook[title] if title in workbook.sheetnames else workbook.create_sheet(title)
    if sheet.max_row:
        sheet.delete_rows(1, sheet.max_row)

    headers = ("Niveau", "Élément", "Tag", "Valeur")
    sheet.append(headers)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    rows: list[tuple[str, str, str, str]] = []
    seen: set[tuple[str, str, str, str]] = set()
    for probe in probes:
        for tag, value in probe.template_tags:
            row = ("Template", probe.template_name, tag, value)
            if row not in seen:
                seen.add(row)
                rows.append(row)
        for trigger in probe.triggers:
            for tag, value in trigger.tags:
                row = ("Trigger", trigger.name, tag, value)
                if row not in seen:
                    seen.add(row)
                    rows.append(row)

    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{max(sheet.max_row, 1)}"
    widths = (14, 45, 28, 35)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width


def generate_catalogue(
    model_path: str | Path,
    output_path: str | Path,
    probes: list[ProbeDefinition],
    *,
    sheet_name: str | None = None,
) -> Path:
    """Populate a copy of the catalogue model with active probes and metadata."""

    output = Path(output_path)
    workbook = load_workbook(Path(model_path))
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    header_row, columns = _find_header_row(sheet)
    first_data_row = header_row + 2 if sheet.cell(header_row + 1, 1).value is None else header_row + 1
    row_style = _capture_row_style(sheet, first_data_row)

    if sheet.max_row >= first_data_row:
        sheet.delete_rows(first_data_row, sheet.max_row - first_data_row + 1)

    output_rows: list[tuple[ProbeDefinition, TriggerDefinition | None]] = []
    for probe in probes:
        if probe.triggers:
            output_rows.extend((probe, trigger) for trigger in probe.triggers)
        else:
            output_rows.append((probe, None))

    for offset, (probe, trigger) in enumerate(output_rows):
        row = first_data_row + offset
        _apply_row_style(sheet, row, row_style)
        for key, value in _row_values(probe, trigger).items():
            _set(sheet, row, columns, key, value)
        if trigger:
            _apply_trigger_style(sheet, row, columns, trigger)

    sheet.auto_filter.ref = f"A{header_row}:{sheet.cell(header_row, sheet.max_column).coordinate}"
    sheet.freeze_panes = f"A{first_data_row}"
    _write_tags_sheet(workbook, probes)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output
