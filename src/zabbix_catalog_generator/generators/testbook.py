"""Build and export test scenarios from parsed Zabbix probes."""

from __future__ import annotations

import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..models import ProbeDefinition, TestCase

TESTBOOK_HEADERS: dict[str, tuple[str, ...]] = {
    "policy_name": ("Nom de la politique", "Politique", "Policy name", "Policy"),
    "resource": ("Ressource", "Resource"),
    "trigger_name": (
        "Nom du déclencheur",
        "Déclencheur",
        "Détection",
        "Trigger name",
        "Trigger",
        "Detection",
    ),
    "severity": ("Sévérité", "Severite", "Severity"),
    "test_type": ("Type de test", "Type", "Test type"),
    "macros": ("Macros", "Macro", "Macros du déclencheur", "Trigger macros"),
    "prerequisites": ("Prérequis", "Prerequis", "Prerequisites"),
    "action": ("Action", "Actions"),
    "expected_result": ("Résultat attendu", "Resultat attendu", "Expected result"),
    "comments": ("Commentaires", "Commentaire", "Comments", "Comment"),
}

# Severity and manual columns are optional because compact testbook models may
# intentionally expose only the columns needed to execute and record a test.
_REQUIRED_HEADERS = {
    "policy_name",
    "resource",
    "trigger_name",
    "test_type",
    "expected_result",
}

_USER_MACRO_PATTERN = re.compile(r"\{\$[^{}]+\}")


def _normalise(value: object) -> str:
    """Normalise an Excel header for resilient matching."""

    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _find_header_row(sheet: Worksheet) -> tuple[int, dict[str, int]]:
    expected = {
        _normalise(label): field
        for field, labels in TESTBOOK_HEADERS.items()
        for label in labels
    }

    for row in range(1, min(sheet.max_row, 50) + 1):
        columns: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            field = expected.get(_normalise(sheet.cell(row, column).value))
            if field:
                columns[field] = column

        if _REQUIRED_HEADERS.issubset(columns):
            return row, columns

    missing = ", ".join(sorted(_REQUIRED_HEADERS))
    raise ValueError(
        "The testbook model does not contain the expected header row "
        f"(required fields: {missing})"
    )


def _capture_row_style(
    sheet: Worksheet,
    source_row: int,
) -> tuple[float | None, list[dict[str, object]]]:
    styles: list[dict[str, object]] = []
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(source_row, column)
        styles.append(
            {
                "style": copy(source._style),
                "number_format": source.number_format,
                "alignment": copy(source.alignment),
                "protection": copy(source.protection),
            }
        )
    return sheet.row_dimensions[source_row].height, styles


def _apply_row_style(
    sheet: Worksheet,
    target_row: int,
    row_style: tuple[float | None, list[dict[str, object]]],
) -> None:
    height, styles = row_style
    sheet.row_dimensions[target_row].height = height
    for column, style in enumerate(styles, start=1):
        target = sheet.cell(target_row, column)
        target._style = copy(style["style"])
        target.number_format = str(style["number_format"])
        target.alignment = copy(style["alignment"])
        target.protection = copy(style["protection"])


def _first_data_row(sheet: Worksheet, header_row: int) -> int:
    """Return the first styled data row, allowing one spacer row below headers."""

    next_row = header_row + 1
    if next_row <= sheet.max_row and all(
        sheet.cell(next_row, column).value is None
        for column in range(1, sheet.max_column + 1)
    ):
        return next_row + 1
    return next_row


def _write_test_case(
    sheet: Worksheet,
    row: int,
    columns: dict[str, int],
    test_case: TestCase,
) -> None:
    for field, column in columns.items():
        sheet.cell(row, column).value = getattr(test_case, field)


def _trigger_macros(probe: ProbeDefinition, expression: str, recovery_expression: str) -> str:
    """Return only user macros referenced by a trigger, preserving expression order."""

    values = dict(probe.template_macros)
    seen: set[str] = set()
    lines: list[str] = []
    for macro in _USER_MACRO_PATTERN.findall(f"{expression}\n{recovery_expression}"):
        if macro in seen:
            continue
        seen.add(macro)
        value = values.get(macro)
        lines.append(f"{macro}={value}" if value is not None else f"{macro}=<non définie>")
    return "\n".join(lines)


def build_test_cases(probes: list[ProbeDefinition]) -> list[TestCase]:
    """Create problem and recovery scenarios for enabled trigger definitions."""

    test_cases: list[TestCase] = []
    for probe in probes:
        for trigger in probe.triggers:
            macros = _trigger_macros(
                probe,
                trigger.expression,
                trigger.recovery_expression,
            )
            test_cases.append(
                TestCase(
                    policy_name=probe.template_name,
                    resource=probe.name,
                    trigger_name=trigger.name,
                    severity=trigger.severity,
                    test_type="PROBLEM",
                    macros=macros,
                    expected_result=f"Le déclencheur « {trigger.name} » passe en état PROBLEM.",
                )
            )
            if trigger.recovery_expression:
                test_cases.append(
                    TestCase(
                        policy_name=probe.template_name,
                        resource=probe.name,
                        trigger_name=trigger.name,
                        severity=trigger.severity,
                        test_type="RECOVERY",
                        macros=macros,
                        expected_result=(
                            f"Le déclencheur « {trigger.name} » revient à l'état OK."
                        ),
                    )
                )
    return test_cases


def generate_testbook(
    model_path: str | Path,
    output_path: str | Path,
    test_cases: list[TestCase],
    *,
    sheet_name: str | None = None,
) -> Path:
    """Populate a copy of the Excel testbook model with generated test cases."""

    output = Path(output_path)
    workbook = load_workbook(Path(model_path))
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    header_row, columns = _find_header_row(sheet)
    first_data_row = _first_data_row(sheet, header_row)
    row_style = _capture_row_style(sheet, first_data_row)

    if sheet.max_row >= first_data_row:
        sheet.delete_rows(first_data_row, sheet.max_row - first_data_row + 1)

    for offset, test_case in enumerate(test_cases):
        row = first_data_row + offset
        _apply_row_style(sheet, row, row_style)
        _write_test_case(sheet, row, columns, test_case)

    sheet.auto_filter.ref = f"A{header_row}:{sheet.cell(header_row, sheet.max_column).coordinate}"
    sheet.freeze_panes = f"A{first_data_row}"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output
