from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from zabbix_catalog_generator.generators import (
    build_test_cases,
    generate_catalogue,
    generate_testbook,
)
from zabbix_catalog_generator.models import ProbeDefinition, TestCase, TriggerDefinition
from zabbix_catalog_generator.scenarios import load_scenarios


def _probe(*triggers: TriggerDefinition) -> ProbeDefinition:
    return ProbeDefinition(
        template_name="Linux by Zabbix agent",
        template_description="",
        name="CPU utilization",
        key="system.cpu.util",
        description="",
        delay="1m",
        item_type="ZABBIX_PASSIVE",
        value_type="FLOAT",
        units="%",
        history="7d",
        trends="365d",
        template_macros=(
            ("{$CPU.MAX}", "90"),
            ("{$CPU.RECOVERY}", "80"),
            ("{$UNUSED}", "123"),
        ),
        triggers=triggers,
    )


def _create_model(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tests"
    headers = [
        "Politique",
        "Ressource",
        "Déclencheur",
        "Sévérité",
        "Expression",
        "Macros",
        "Tags",
        "Code test",
        "Scénario",
        "Objectif",
        "Prérequis",
        "Procédure",
        "Résultat attendu",
        "Testeur",
        "Date",
        "Statut",
        "Preuve",
        "Commentaires",
        "Génération",
    ]
    sheet.append(headers)
    sheet.append([None] * len(headers))
    sheet.append([None] * len(headers))
    for cell in sheet[3]:
        cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        cell.font = Font(name="Arial", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[3].height = 42
    workbook.save(path)


def test_test_case_qualification_fields_are_empty_by_default() -> None:
    test_case = TestCase(
        policy_name="Template",
        resource="Resource",
        trigger_name="Trigger",
        severity="HIGH",
    )

    assert test_case.expression == ""
    assert test_case.macros == ""
    assert test_case.tags == ""
    assert test_case.test_code == ""
    assert test_case.procedure == ""
    assert test_case.tester == ""
    assert test_case.evidence == ""
    assert test_case.generation == "AUTO"


def test_default_repository_contains_nominal_scenario() -> None:
    scenarios = load_scenarios()

    assert [scenario.code for scenario in scenarios] == ["TEST-TRG-001"]
    assert scenarios[0].applies == "always"


def test_build_test_cases_renders_nominal_scenario_and_trigger_variables() -> None:
    trigger = TriggerDefinition(
        name="CPU usage is high",
        severity="HIGH",
        expression="last(/template/system.cpu.util)>{$CPU.MAX}",
        recovery_expression="last(/template/system.cpu.util)<{$CPU.RECOVERY}",
        tags=(("application", "linux"), ("team", "production")),
    )

    test_cases = build_test_cases([_probe(trigger)])

    assert len(test_cases) == 1
    test_case = test_cases[0]
    assert test_case.test_code == "TEST-TRG-001"
    assert test_case.scenario == "Déclenchement nominal"
    assert test_case.policy_name == "Linux by Zabbix agent"
    assert test_case.resource == "CPU utilization"
    assert test_case.trigger_name == "CPU usage is high"
    assert test_case.severity == "HIGH"
    assert test_case.expression == trigger.expression
    assert test_case.macros == "{$CPU.MAX}=90\n{$CPU.RECOVERY}=80"
    assert test_case.tags == "application=linux\nteam=production"
    assert "CPU usage is high" in test_case.objective
    assert "CPU utilization" in test_case.prerequisites
    assert "{$CPU.MAX}=90" in test_case.procedure
    assert "1m" in test_case.procedure
    assert "${" not in test_case.procedure


def test_build_test_cases_only_exports_macros_used_by_trigger() -> None:
    trigger = TriggerDefinition(
        name="CPU usage is high",
        expression=(
            "avg(/template/system.cpu.util,5m)>{$CPU.MAX} "
            "and last(/template/system.cpu.util)>{$CPU.MAX}"
        ),
    )

    test_case = build_test_cases([_probe(trigger)])[0]

    assert test_case.macros == "{$CPU.MAX}=90"
    assert "{$UNUSED}" not in test_case.macros


def test_build_test_cases_marks_undefined_trigger_macro() -> None:
    trigger = TriggerDefinition(
        name="CPU usage is high",
        expression="last(/template/system.cpu.util)>{$CPU.UNKNOWN}",
    )

    test_case = build_test_cases([_probe(trigger)])[0]

    assert test_case.macros == "{$CPU.UNKNOWN}=<non définie>"


def test_custom_scenario_repository_supports_conditional_macros(tmp_path: Path) -> None:
    repository = tmp_path / "scenarios.yaml"
    repository.write_text(
        """
- code: TEST-TRG-003
  title: Validation du seuil
  applies: uses_macros
  objective: Vérifier les macros de ${trigger.name}.
  procedure: |-
    Contrôler les macros :
    ${trigger.macros}
  expected_result: Le seuil est conforme.
""".strip(),
        encoding="utf-8",
    )

    without_macro = TriggerDefinition(name="No macro", expression="last(/t/key)>90")
    with_macro = TriggerDefinition(name="With macro", expression="last(/t/key)>{$CPU.MAX}")

    test_cases = build_test_cases(
        [_probe(without_macro, with_macro)],
        scenarios_path=repository,
    )

    assert [test_case.trigger_name for test_case in test_cases] == ["With macro"]
    assert test_cases[0].test_code == "TEST-TRG-003"
    assert "{$CPU.MAX}=90" in test_cases[0].procedure


def test_loader_rejects_duplicate_scenario_codes(tmp_path: Path) -> None:
    repository = tmp_path / "scenarios.yaml"
    repository.write_text(
        """
- code: TEST-TRG-001
  title: First
  objective: First objective
  procedure: First procedure
  expected_result: First result
- code: TEST-TRG-001
  title: Second
  objective: Second objective
  procedure: Second procedure
  expected_result: Second result
""".strip(),
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="Duplicate scenario codes"):
        load_scenarios(repository)


def test_generators_package_keeps_catalogue_api() -> None:
    assert callable(generate_catalogue)


def test_generate_testbook_populates_v2_model_and_preserves_style(tmp_path: Path) -> None:
    model = tmp_path / "model.xlsx"
    output = tmp_path / "output.xlsx"
    _create_model(model)
    test_cases = [
        TestCase(
            policy_name="Linux by Zabbix agent",
            resource="CPU utilization",
            trigger_name="CPU usage is high",
            severity="HIGH",
            expression="last(/template/system.cpu.util)>{$CPU.MAX}",
            macros="{$CPU.MAX}=90",
            tags="application=linux",
            test_code="TEST-TRG-001",
            scenario="Déclenchement nominal",
            objective="Vérifier le déclenchement.",
            prerequisites="Agent actif",
            procedure="Générer une charge CPU",
            expected_result="Le déclencheur passe en état PROBLÈME.",
            comments="Test automatique",
        )
    ]

    result = generate_testbook(model, output, test_cases)

    assert result == output
    workbook = load_workbook(output)
    sheet = workbook["Tests"]
    assert sheet.cell(3, 1).value == "Linux by Zabbix agent"
    assert sheet.cell(3, 5).value == "last(/template/system.cpu.util)>{$CPU.MAX}"
    assert sheet.cell(3, 6).value == "{$CPU.MAX}=90"
    assert sheet.cell(3, 8).value == "TEST-TRG-001"
    assert sheet.cell(3, 9).value == "Déclenchement nominal"
    assert sheet.cell(3, 12).value == "Générer une charge CPU"
    assert sheet.cell(3, 19).value == "AUTO"
    assert sheet.cell(3, 1).fill.fgColor.rgb == sheet.cell(3, 2).fill.fgColor.rgb
    assert sheet.cell(3, 1).font.bold is True
    assert sheet.row_dimensions[3].height == 42
    assert sheet.freeze_panes == "A3"


def test_generate_testbook_rejects_model_without_expected_headers(tmp_path: Path) -> None:
    model = tmp_path / "invalid.xlsx"
    workbook = Workbook()
    workbook.active.append(["Unknown", "Columns"])
    workbook.save(model)

    with pytest.raises(ValueError, match="expected header row"):
        generate_testbook(model, tmp_path / "output.xlsx", [])
