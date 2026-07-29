from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from zabbix_catalog_generator.cli import main


def _write_template(path: Path) -> None:
    path.write_text(
        """
zabbix_export:
  templates:
    - name: Demo template
      items:
        - name: CPU utilization
          key: system.cpu.util
          delay: 1m
          triggers:
            - name: CPU usage is high
              expression: last(/Demo template/system.cpu.util)>90
              recovery_expression: last(/Demo template/system.cpu.util)<80
              priority: HIGH
""",
        encoding="utf-8",
    )


def _write_testbook_model(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tests"
    sheet.append(
        [
            "Politique",
            "Ressource",
            "Déclencheur",
            "Type",
            "Action",
            "Résultat attendu",
            "Statut",
        ]
    )
    sheet.append([None] * 7)
    workbook.save(path)


def test_tests_command_generates_problem_and_recovery_rows(
    tmp_path: Path,
    capsys,
) -> None:
    template = tmp_path / "template.yaml"
    model = tmp_path / "Modele_Test.xlsx"
    output = tmp_path / "Cahier_Test.xlsx"
    _write_template(template)
    _write_testbook_model(model)

    exit_code = main(
        [
            "tests",
            str(template),
            "--model",
            str(model),
            "--output",
            str(output),
        ]
    )

    assert exit_code == 0
    assert output.exists()

    workbook = load_workbook(output)
    sheet = workbook["Tests"]
    assert sheet.cell(2, 1).value == "Demo template"
    assert sheet.cell(2, 2).value == "CPU utilization"
    assert sheet.cell(2, 3).value == "CPU usage is high"
    assert sheet.cell(2, 4).value == "PROBLEM"
    assert sheet.cell(2, 5).value is None
    assert "PROBLEM" in sheet.cell(2, 6).value
    assert sheet.cell(2, 7).value is None

    assert sheet.cell(3, 1).value == "Demo template"
    assert sheet.cell(3, 3).value == "CPU usage is high"
    assert sheet.cell(3, 4).value == "RECOVERY"
    assert "OK" in sheet.cell(3, 6).value

    output_text = capsys.readouterr().out
    assert "Cahier de tests généré (2 scénarios)" in output_text


def test_tests_command_returns_one_for_invalid_model(
    tmp_path: Path,
    capsys,
) -> None:
    template = tmp_path / "template.yaml"
    model = tmp_path / "invalid.xlsx"
    output = tmp_path / "Cahier_Test.xlsx"
    _write_template(template)

    workbook = Workbook()
    workbook.active.append(["Colonne inconnue"])
    workbook.save(model)

    exit_code = main(
        [
            "tests",
            str(template),
            "--model",
            str(model),
            "--output",
            str(output),
        ]
    )

    assert exit_code == 1
    assert not output.exists()
    assert "expected header row" in capsys.readouterr().err
