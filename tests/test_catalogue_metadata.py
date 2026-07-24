from pathlib import Path

from openpyxl import Workbook, load_workbook

from zabbix_catalog_generator.excel import generate_catalogue
from zabbix_catalog_generator.models import ProbeDefinition, TriggerDefinition


def test_catalogue_writes_contextual_metadata_and_tags(tmp_path: Path) -> None:
    model = tmp_path / "model.xlsx"
    output = tmp_path / "catalogue.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Catalogue"
    sheet.append(
        [
            "Ressource",
            "Clef",
            "Description",
            "Prétraitement",
            "Filtres",
            "Macros",
            "Triger Name",
        ]
    )
    sheet.append([None] * 7)
    sheet.append(["style"] * 7)
    workbook.save(model)

    trigger = TriggerDefinition(
        name="CPU high",
        expression="last(/Demo/demo.cpu)>{$CPU.MAX}",
        tags=(("scope", "performance"),),
    )
    probe = ProbeDefinition(
        template_name="Demo",
        template_description="",
        name="CPU {$CPU.MAX}",
        key="demo.cpu",
        description="CPU usage",
        delay="1m",
        item_type="ZABBIX_ACTIVE",
        value_type="FLOAT",
        units="%",
        history="7d",
        trends="365d",
        preprocessing=("JSONPATH: $.cpu", "MULTIPLIER: 100"),
        filters=("A: {#CPU} matches regex ^cpu",),
        template_macros=(("{$CPU.MAX}", "90"), ("{$UNUSED}", "x")),
        template_tags=(("component", "system"),),
        triggers=(trigger,),
    )

    generate_catalogue(model, output, [probe])

    result = load_workbook(output)
    catalogue = result["Catalogue"]
    assert catalogue.cell(3, 4).value == "JSONPATH: $.cpu\nMULTIPLIER: 100"
    assert catalogue.cell(3, 5).value == "A: {#CPU} matches regex ^cpu"
    assert catalogue.cell(3, 6).value == "{$CPU.MAX} = 90"

    tags = result["Tags"]
    assert [cell.value for cell in tags[1]] == ["Niveau", "Élément", "Tag", "Valeur"]
    assert [cell.value for cell in tags[2]] == ["Template", "Demo", "component", "system"]
    assert [cell.value for cell in tags[3]] == ["Trigger", "CPU high", "scope", "performance"]
