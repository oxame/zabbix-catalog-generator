from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText

from zabbix_catalog_generator.excel import generate_catalogue
from zabbix_catalog_generator.models import ProbeDefinition, TriggerDefinition


def _probe(*, triggers: tuple[TriggerDefinition, ...] = ()) -> ProbeDefinition:
    return ProbeDefinition(
        template_name="Oracle",
        template_description="",
        name="Sessions",
        key="oracle.sessions",
        description="Session count",
        delay="5m",
        item_type="ZABBIX_ACTIVE",
        value_type="UNSIGNED",
        units="",
        history="7d",
        trends="365d",
        triggers=triggers,
    )


def test_catalogue_creates_one_row_per_trigger(tmp_path: Path) -> None:
    model = tmp_path / "model.xlsx"
    output = tmp_path / "catalogue.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Database"
    sheet.append(["Ressource", "Clef", "Description", "LLD", "Triger Name", "Sévérité"])
    sheet.append([None] * 6)
    sheet.append(["style"] * 6)
    workbook.save(model)

    probe = _probe(
        triggers=(
            TriggerDefinition("Warning", severity="WARNING"),
            TriggerDefinition("High", severity="HIGH"),
        )
    )

    generate_catalogue(model, output, [probe])

    result = load_workbook(output).active
    assert result.cell(3, 1).value == "Sessions"
    assert result.cell(4, 1).value == "Sessions"
    assert result.cell(3, 5).value == "Warning"
    assert result.cell(4, 5).value == "High"


def test_catalogue_colors_trigger_name_and_severity(tmp_path: Path) -> None:
    model = tmp_path / "model.xlsx"
    output = tmp_path / "catalogue.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Ressource", "Clef", "Description", "Triger Name", "Sévérité"])
    sheet.append([None] * 5)
    sheet.append(["style"] * 5)
    workbook.save(model)

    severities = (
        ("NOT_CLASSIFIED", "97AAB3", "FFFFFF"),
        ("INFORMATION", "7499FF", "FFFFFF"),
        ("WARNING", "FFC859", "000000"),
        ("AVERAGE", "FFA059", "000000"),
        ("HIGH", "E97659", "FFFFFF"),
        ("DISASTER", "E45959", "FFFFFF"),
    )
    probe = _probe(
        triggers=tuple(
            TriggerDefinition(f"Trigger {severity}", severity=severity)
            for severity, _, _ in severities
        )
    )

    generate_catalogue(model, output, [probe])

    result = load_workbook(output).active
    for row, (_, fill, font) in enumerate(severities, start=3):
        for column in (4, 5):
            cell = result.cell(row, column)
            assert cell.fill.fill_type == "solid"
            assert cell.fill.fgColor.rgb == f"00{fill}"
            assert cell.font.color.rgb == f"00{font}"


def test_catalogue_keeps_default_style_without_trigger(tmp_path: Path) -> None:
    model = tmp_path / "model.xlsx"
    output = tmp_path / "catalogue.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Ressource", "Clef", "Description", "Triger Name", "Sévérité"])
    sheet.append([None] * 5)
    sheet.append(["style"] * 5)
    workbook.save(model)

    generate_catalogue(model, output, [_probe()])

    result = load_workbook(output).active
    for column in (4, 5):
        cell = result.cell(3, column)
        assert cell.fill.fill_type is None


def test_catalogue_highlights_dependency_before_probe_name(tmp_path: Path) -> None:
    model = tmp_path / "model.xlsx"
    output = tmp_path / "catalogue.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Ressource", "Clef", "Description"])
    sheet.append([None] * 3)
    sheet.append(["style"] * 3)
    workbook.save(model)

    probe = ProbeDefinition(
        template_name="Oracle",
        template_description="",
        name="Sessions used percentage",
        key="oracle.sessions.pct",
        description="Calculated percentage",
        delay="5m",
        item_type="CALCULATED",
        value_type="FLOAT",
        units="%",
        history="7d",
        trends="365d",
        dependency_names=("Sessions current", "Sessions limit"),
    )

    generate_catalogue(model, output, [probe])

    result = load_workbook(output, rich_text=True).active
    resource = result.cell(3, 1).value
    assert isinstance(resource, CellRichText)
    assert str(resource) == (
        "Dépend de : Sessions current\n"
        "Dépend de : Sessions limit\n"
        "Sessions used percentage"
    )
    assert resource[0].font.color.rgb == "00008000"
